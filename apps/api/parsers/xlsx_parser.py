from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def _serializable(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def parse_xlsx(path: str | Path) -> list[dict]:
    workbook = load_workbook(filename=str(path), data_only=False, read_only=False)
    value_workbook = load_workbook(filename=str(path), data_only=True, read_only=False)
    blocks: list[dict] = []
    for sheet in workbook.worksheets:
        merged = [str(cell_range) for cell_range in sheet.merged_cells.ranges]
        headers = [_serializable(sheet.cell(1, col).value) for col in range(1, sheet.max_column + 1)] if sheet.max_row else []
        for row_number in range(1, sheet.max_row + 1):
            cells: list[dict[str, Any]] = []
            display: list[str] = []
            for column in range(1, sheet.max_column + 1):
                cell = sheet.cell(row_number, column)
                if cell.value is None:
                    continue
                address = f"{get_column_letter(column)}{row_number}"
                value = _serializable(cell.value)
                calculated_value = _serializable(value_workbook[sheet.title][address].value) if cell.data_type == "f" else value
                cells.append(
                    {
                        "address": address,
                        "value": value,
                        "header": headers[column - 1] if row_number > 1 and column - 1 < len(headers) else None,
                        "data_type": cell.data_type,
                        "is_formula": cell.data_type == "f",
                        "calculated_value": calculated_value,
                    }
                )
                display.append(f"{address}={value}")
            if not cells:
                continue
            blocks.append(
                {
                    "block_type": "spreadsheet_row",
                    "content": " | ".join(display),
                    "block_order": len(blocks),
                    "page_number": None,
                    "sheet_name": sheet.title,
                    "section_title": None,
                    "location_metadata": {
                        "row_number": row_number,
                        "cells": cells,
                        "headers": headers,
                        "merged_cells": merged,
                    },
                }
            )
    if not blocks:
        raise ValueError("XLSX에서 값이 있는 셀을 찾지 못했습니다")
    return blocks
